"""Course notes loader and rolling transcript buffer."""
from __future__ import annotations

import time
from collections import deque
from dataclasses import dataclass
from pathlib import Path


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".tex", ".csv", ".py"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".pptx":
        from pptx import Presentation
        slides = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
            slides.append(f"[Slide {i}] " + " | ".join(t for t in texts if t.strip()))
        return "\n".join(slides)
    if suffix == ".docx":
        import zipfile
        import re
        xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8", "ignore")
        return re.sub(r"<[^>]+>", " ", xml.replace("</w:p>", "\n"))
    return ""


def load_notes(paths, max_chars=60000) -> tuple[str, list[str]]:
    """Concatenate notes from files/folders; returns (text, list of files used)."""
    files = []
    for p in (Path(x).expanduser() for x in paths):
        if p.is_dir():
            files += sorted(f for f in p.rglob("*") if f.is_file())
        elif p.is_file():
            files.append(p)
    chunks, used, total = [], [], 0
    for f in files:
        try:
            text = read_document(f).strip()
        except Exception as error:   # unreadable file: skip, keep going
            text = ""
            print(f"warning: could not read {f}: {error}")
        if not text:
            continue
        block = f"### {f.name}\n{text}\n"
        if total + len(block) > max_chars:
            block = block[: max(0, max_chars - total)]
        chunks.append(block)
        used.append(str(f))
        total += len(block)
        if total >= max_chars:
            break
    return "\n".join(chunks), used


@dataclass
class Line:
    t: float
    text: str


class TranscriptBuffer:
    def __init__(self):
        self.lines: deque[Line] = deque(maxlen=5000)

    def add(self, text: str, t: float | None = None):
        self.lines.append(Line(time.time() if t is None else t, text))

    def window(self, minutes: float) -> str:
        cutoff = time.time() - minutes * 60
        return "\n".join(l.text for l in self.lines if l.t >= cutoff)

    def last(self, n=1) -> str:
        return " ".join(l.text for l in list(self.lines)[-n:])
