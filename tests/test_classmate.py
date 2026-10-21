import asyncio
import json
import time
import wave
from urllib.parse import parse_qs, urlparse

import numpy as np
import pytest
from aiohttp.test_utils import TestClient, TestServer
from websockets.asyncio.server import serve

from classmate.app import Session, build_app
from classmate.asr import Transcript
from classmate.asr.assemblyai import AssemblyAI
from classmate.asr.deepgram import Deepgram
from classmate.asr.elevenlabs import ElevenLabs
from classmate.audio import FileSource
from classmate.config import Settings, load_settings
from classmate.context import TranscriptBuffer, load_notes
from classmate.llm import build_prompt
from classmate.questions import find_question, question_score
from classmate import license as licensing

REAL_LICENSE_LOAD = licensing.load


@pytest.fixture(autouse=True)
def isolated_config(tmp_path, monkeypatch):
    """Keep the real .env, settings.yaml and API keys of this machine out of the tests."""
    from classmate import config
    home = tmp_path / "app"
    home.mkdir()
    monkeypatch.setattr(config, "APP_DIR", home)
    monkeypatch.setattr(config, "UI_SETTINGS", home / "settings.yaml")
    monkeypatch.setattr(config, "ENV_FILE", None)
    monkeypatch.setattr(config, "NOTES_DIR", home / "notes")
    monkeypatch.chdir(home)
    for env in set(Settings.API_ENV.values()):
        monkeypatch.delenv(env, raising=False)
    from classmate import license as licensing     # most tests run as an authorized user
    monkeypatch.setattr(licensing, "load", lambda now=None: licensing.License("user", None, "test"))


# ---------------- question detection ----------------
@pytest.mark.parametrize("text", [
    "What is the difference between a homography and an affine transform?",
    "So why do we normalise the features before the dot product",
    "Can anyone tell me what the receptive field is",
    "Does anyone know why max pooling helps here?",
    "Okay, how would you choose the learning rate",
    "这个卷积核的作用是什么？",
    "大家觉得为什么要做数据增强",
    "What do you think happens if we remove the bias term",
    "How does the receptive field grow with depth?",
    "有没有人知道反向传播的原理",
])
def test_detects_questions(text):
    assert find_question(text) is not None, question_score(text)


@pytest.mark.parametrize("text", [
    "Today we will cover convolutional networks.",
    "This is the key idea behind backpropagation.",
    "Okay.",
    "我们今天讲卷积神经网络。",
    "So the answer is that the gradient vanishes.",
    "What we'll do today is look at feature matching.",
    "How this works is that each filter slides over the image.",
    "Which means the loss goes down.",
])
def test_ignores_statements(text):
    assert find_question(text) is None, question_score(text)


def test_extracts_question_inside_longer_turn():
    turn = "We looked at SIFT last week. What makes it invariant to rotation? Think about the orientation histogram."
    q = find_question(turn)
    assert "invariant to rotation?" in q and "Think about" not in q


# ---------------- provider message parsing ----------------
def settings(**kw):
    return Settings(**kw)


def test_assemblyai_parse_and_url(monkeypatch):
    monkeypatch.setenv("ASSEMBLYAI_API_KEY", "k")
    asr = AssemblyAI(settings(language="zh", keyterms=["homography"]))
    q = parse_qs(urlparse(asr.build_url()).query)
    assert q["speech_model"] == ["universal-3-6-pro"] and q["language_codes"] == ["zh"]
    assert q["keyterms_prompt"] == ["homography"] and q["encoding"] == ["pcm_s16le"]
    assert asr.headers() == {"Authorization": "k"}
    assert asr.parse({"type": "Begin"}) == []
    [e] = asr.parse({"type": "Turn", "transcript": "what is", "end_of_turn": False})
    assert not e.final
    [e] = asr.parse({"type": "Turn", "transcript": "what is a kernel", "end_of_turn": True})
    assert e.final and e.end_of_turn
    legacy = AssemblyAI(settings(asr_model="universal-streaming-english"))
    assert "format_turns=true" in legacy.build_url()
    [e] = legacy.parse({"type": "Turn", "transcript": "what is a kernel", "end_of_turn": True,
                        "turn_is_formatted": False})
    assert not e.final


def test_elevenlabs_parse(monkeypatch):
    monkeypatch.setenv("ELEVENLABS_API_KEY", "k")
    asr = ElevenLabs(settings())
    assert "language_code" not in asr.build_url() and "pcm_16000" in asr.build_url()
    assert asr.parse({"message_type": "partial_transcript", "text": "why"})[0].final is False
    assert asr.parse({"message_type": "committed_transcript", "text": "why so?"})[0].end_of_turn
    with pytest.raises(RuntimeError):
        asr.parse({"message_type": "rate_limited", "error": "slow down"})


def test_deepgram_assembles_segments(monkeypatch):
    monkeypatch.setenv("DEEPGRAM_API_KEY", "k")
    asr = Deepgram(settings())
    assert asr.headers()["Authorization"] == "Token k"
    r = lambda text, final, speech_final=False: {"type": "Results", "is_final": final, "speech_final": speech_final,
                                                 "channel": {"alternatives": [{"transcript": text}]}}
    assert asr.parse(r("what is", False))[0].text == "what is"
    assert not asr.parse(r("what is the", True))[0].final
    [e] = asr.parse(r("receptive field", True, True))
    assert e.final and e.text == "what is the receptive field"
    asr.parse(r("any ideas", True))
    [e] = asr.parse({"type": "UtteranceEnd"})
    assert e.end_of_turn and e.text == "any ideas"


async def test_streaming_loop_against_fake_server(monkeypatch):
    """Full websocket loop: audio frames out, Turn events in, Terminate at the end."""
    monkeypatch.setenv("ASSEMBLYAI_API_KEY", "secret")
    received = {"bytes": 0, "terminate": False, "auth": None}

    async def handler(ws):
        received["auth"] = ws.request.headers.get("Authorization")
        await ws.send(json.dumps({"type": "Begin", "id": "x"}))
        async for msg in ws:
            if isinstance(msg, bytes):
                received["bytes"] += len(msg)
                if received["bytes"] == 3200:
                    await ws.send(json.dumps({"type": "Turn", "transcript": "what is", "end_of_turn": False}))
            elif json.loads(msg)["type"] == "Terminate":
                received["terminate"] = True
                await ws.send(json.dumps({"type": "Turn", "transcript": "What is a kernel?", "end_of_turn": True}))
                await ws.send(json.dumps({"type": "Termination"}))
                await ws.close()

    async with serve(handler, "127.0.0.1", 0) as server:
        port = server.sockets[0].getsockname()[1]
        asr = AssemblyAI(settings())
        asr.url_override = f"ws://127.0.0.1:{port}"
        queue = asyncio.Queue()
        for _ in range(4):
            queue.put_nowait(bytes(1600))
        queue.put_nowait(None)
        events = []

        async def emit(e):
            events.append(e)
        await asyncio.wait_for(asr.run(queue, emit), 5)
    assert received == {"bytes": 6400, "terminate": True, "auth": "secret"}
    assert [e.final for e in events] == [False, True] and events[-1].text == "What is a kernel?"


# ---------------- context ----------------
def test_notes_loader_and_prompt(tmp_path):
    (tmp_path / "w10.md").write_text("Homography maps planes between views.")
    from pptx import Presentation
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "RANSAC"
    deck.save(tmp_path / "slides.pptx")
    text, files = load_notes([str(tmp_path)], max_chars=10000)
    assert "Homography" in text and "RANSAC" in text and len(files) == 2
    blocks, user = build_prompt("What is RANSAC?", "transcript here", text)
    assert "course_notes" in blocks[-1] and "What is RANSAC?" in user and "transcript here" in user
    clipped, _ = load_notes([str(tmp_path)], max_chars=30)
    assert len(clipped) <= 31


def test_transcript_window():
    buf = TranscriptBuffer()
    buf.add("old", t=0)
    buf.add("new")
    assert buf.window(1) == "new" and buf.last(2) == "old new"


def test_settings_reject_unknown_keys(tmp_path):
    (tmp_path / "c.yaml").write_text("asr: deepgram\nbogus: 1\n")
    with pytest.raises(ValueError):
        load_settings(str(tmp_path / "c.yaml"))
    (tmp_path / "c.yaml").write_text("asr: deepgram\n")
    assert load_settings(str(tmp_path / "c.yaml"), language="zh").language == "zh"


# ---------------- end to end: audio file -> ASR -> question -> answer -> UI ----------------
class ScriptedASR:
    """Consumes the real audio stream and emits a scripted lecture."""
    script = [Transcript("today we cover convolution", False),
              Transcript("Today we cover convolution.", True, True),
              Transcript("why do we use", False),
              Transcript("Why do we use padding in a convolution layer?", True, True)]

    async def run(self, queue, emit):
        frames = 0
        while (chunk := await queue.get()) is not None:
            frames += 1
            if frames % 3 == 0 and self.script:
                await emit(self.script.pop(0))
        self.frames = frames


class FakeAnswerer:
    def __init__(self):
        self.calls = []

    async def stream(self, question, transcript, notes, deep=False, history=(), images=()):
        self.calls.append((question, transcript, deep, list(history)))
        for token in ["Padding keeps ", "the output size", "\n- preserves borders"]:
            await asyncio.sleep(0.01)
            yield token

    async def warmup(self):
        return None


async def test_end_to_end_session(tmp_path):
    wav = tmp_path / "lecture.wav"
    with wave.open(str(wav), "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(16000)
        w.writeframes((np.sin(np.arange(16000) / 5) * 3000).astype(np.int16).tobytes())
    s = Settings(session_dir=str(tmp_path / "sessions"), capture_mode="questions")
    asr, answerer = ScriptedASR(), FakeAnswerer()
    session = Session(s, lambda: FileSource(wav, speed=20), asr=asr, answerer=answerer)
    client = TestClient(TestServer(build_app(session)))
    await client.start_server()
    ws = await client.ws_connect("/ws")
    events = []

    async def collect():
        async for msg in ws:
            events.append(json.loads(msg.data))
            if events[-1]["type"] == "answer_done":
                return
    await ws.send_str(json.dumps({"cmd": "start"}))
    await asyncio.wait_for(collect(), 10)
    kinds = [e["type"] for e in events]
    assert "partial" in kinds and "final" in kinds and "question" in kinds
    q = next(e for e in events if e["type"] == "question")
    assert q["text"] == "Why do we use padding in a convolution layer?" and q["source"] == "auto"
    assert "".join(e["text"] for e in events if e["type"] == "answer_delta").startswith("Padding keeps")
    assert answerer.calls[0][1].startswith("Today we cover convolution.")   # transcript context passed
    done = next(e for e in events if e["type"] == "answer_done")
    assert done["first_token_ms"] >= 0
    # manual triggers and deep re-answer
    await ws.send_str(json.dumps({"cmd": "ask", "text": "What is stride?"}))
    await ws.send_str(json.dumps({"cmd": "deep", "id": q["id"]}))
    await asyncio.sleep(0.5)
    assert ("What is stride?", ) == answerer.calls[1][:1] and answerer.calls[2][2] is True
    replay = session.replay()
    assert [e["type"] for e in replay[:3]] == ["question", "answer_delta", "answer_done"]
    log = session.log_path.read_text()
    assert "Why do we use padding" in log and "Padding keeps the output size" in log
    await ws.close()
    await session.stop()
    await client.close()


async def test_duplicate_auto_questions_are_suppressed(tmp_path):
    session = Session(Settings(session_dir=str(tmp_path)), lambda: None, asr=ScriptedASR(), answerer=FakeAnswerer())
    assert session.ask("What is a kernel?", source="auto") is not None
    assert session.ask("what is a  kernel?", source="auto") is None
    assert session.ask("What is a kernel?", source="manual") is not None
    await asyncio.sleep(0.2)


# ---------------- capture rules (Ctrl hold / extend / queue) and answer history ----------------
class GatedAnswerer:
    """Answers stream only when the test releases the gate, so 'still generating' can be simulated."""
    def __init__(self):
        self.calls, self.gate = [], asyncio.Event()

    async def stream(self, question, transcript, notes, deep=False, history=(), images=()):
        self.calls.append({"q": question, "history": list(history), "images": list(images)})
        await self.gate.wait()
        yield f"answer to: {question}"


async def settle():
    for _ in range(5):
        await asyncio.sleep(0.01)


def make_session(tmp_path, mode="every_turn", style="hold"):
    s = Settings(session_dir=str(tmp_path), capture_mode=mode, capture_style=style)
    answerer = GatedAnswerer()
    return Session(s, lambda: None, asr=ScriptedASR(), answerer=answerer), answerer


async def say(session, text, partial=False):
    await session.on_transcript(Transcript(text, final=not partial, end_of_turn=not partial))


async def test_hold_ctrl_collects_until_release(tmp_path):
    session, ans = make_session(tmp_path, mode="questions")
    await session.ctrl(True)
    await say(session, "Consider a network with three layers.")
    await say(session, "If we double the learning rate, what happens to training?")
    await settle()
    assert ans.calls == []                                   # nothing answered while held
    await session.ctrl(False)
    await settle()
    assert ans.calls[0]["q"] == ("Consider a network with three layers. "
                                 "If we double the learning rate, what happens to training?")


async def test_release_mid_sentence_waits_for_the_sentence_to_end(tmp_path):
    session, ans = make_session(tmp_path)
    await session.ctrl(True)
    await say(session, "So what is the role of")
    await say(session, "the role of the", partial=True)
    await session.ctrl(False)
    await settle()
    assert ans.calls == [] and session.capture_event()["armed"]
    await say(session, "the activation function here?")
    await settle()
    assert ans.calls[0]["q"] == "So what is the role of the activation function here?"


async def test_auto_answer_on_pause_and_queue_next_question(tmp_path):
    session, ans = make_session(tmp_path, mode="every_turn")
    await say(session, "Why do we normalise inputs?")
    await settle()
    assert len(ans.calls) == 1 and session.generating()
    await say(session, "And what about batch norm?")          # arrives while answer 1 streams
    await settle()
    assert len(ans.calls) == 1 and session.pending == ["And what about batch norm?"]
    ans.gate.set()
    await settle()
    assert [c["q"] for c in ans.calls] == ["Why do we normalise inputs?", "And what about batch norm?"]
    assert ans.calls[1]["history"] == [("Why do we normalise inputs?", "answer to: Why do we normalise inputs?")]


async def test_ctrl_after_answer_extends_the_current_question(tmp_path):
    session, ans = make_session(tmp_path)
    ans.gate.set()
    await say(session, "What is dropout?")
    await settle()
    first = session.last_qid
    events = []
    orig = session.broadcast

    async def capture(e):
        events.append(e)
        await orig(e)
    session.broadcast = capture
    await session.ctrl(True)                                  # extend instead of a new question
    await say(session, "And why does it reduce overfitting?")
    await session.ctrl(False)
    await settle()
    assert session.last_qid == first and len(ans.calls) == 2
    assert ans.calls[1]["q"] == "What is dropout? And why does it reduce overfitting?"
    q_event = [e for e in events if e["type"] == "question"][-1]
    assert q_event["id"] == first and q_event["replace"] is True
    assert ans.calls[1]["history"] == []                      # its own old answer is not "earlier"


async def test_ctrl_tap_arms_append_of_next_sentence(tmp_path):
    session, ans = make_session(tmp_path)
    ans.gate.set()
    await say(session, "Define stride.")
    await settle()
    await session.ctrl(True)
    await session.ctrl(False)                                 # tap: nothing said during the press
    assert session.capture_event()["armed"]
    await say(session, "Give an example with a 3 by 3 kernel.")
    await settle()
    assert ans.calls[-1]["q"] == "Define stride. Give an example with a 3 by 3 kernel."


async def test_ctrl_during_generation_appends_queued_speech(tmp_path):
    session, ans = make_session(tmp_path)
    await say(session, "What is a residual connection?")
    await settle()
    await say(session, "Specifically in ResNet-50.")          # queued while generating
    await session.ctrl(True)
    await session.ctrl(False)
    await settle()
    assert ans.calls[-1]["q"] == "What is a residual connection? Specifically in ResNet-50."
    assert session.pending == []


async def test_escape_discards_capture(tmp_path):
    session, ans = make_session(tmp_path)
    await session.ctrl(True)
    await say(session, "half a thought")
    await session.cancel_capture()
    await session.ctrl(False)
    await settle()
    assert ans.calls == [] and session.capture_event()["text"] == ""


def test_prompt_includes_earlier_answers():
    blocks, user = build_prompt("And for colour images?", "", "", history=[("What is a kernel?", "A small filter.")])
    assert "<earlier_questions_and_answers>" in user and "A small filter." in user
    assert user.index("earlier_questions") < user.index("The lecturer just asked")
    _, plain = build_prompt("Q?", "", "")
    assert "earlier_questions" not in plain


def test_window_api_moves_relative():
    from classmate.desktop import WindowAPI

    class W:
        x, y = 100, 200
        def move(self, x, y):
            self.x, self.y = x, y
    api = WindowAPI()
    api.window = W()
    api.move(40, -40)
    assert (api.window.x, api.window.y) == (140, 160)
    api.move(-500, 0)
    assert api.window.x == 0


async def test_default_mode_answers_every_turn(tmp_path):
    session, ans = make_session(tmp_path, mode=Settings().capture_mode)
    assert session.mode == "every_turn"
    await say(session, "Today we cover convolution.")
    await settle()
    assert ans.calls[0]["q"] == "Today we cover convolution."


async def test_deepseek_provider_uses_openai_client_without_thinking(monkeypatch, tmp_path):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "ds-key")
    (tmp_path / "c.yaml").write_text("llm: deepseek\n")
    s = load_settings(str(tmp_path / "c.yaml"))
    assert (s.llm_model, s.llm_deep_model, s.llm_base_url) == ("deepseek-flash", "deepseek-v4-pro",
                                                              "https://api.deepseek.com")
    from classmate.llm import Answerer
    seen = []

    class Chunk:
        def __init__(self, text):
            self.choices = [type("C", (), {"delta": type("D", (), {"content": text})()})()]

    class Completions:
        async def create(self, **kw):
            seen.append(kw)
            async def gen():
                for t in ["Four ", "points."]:
                    yield Chunk(t)
            return gen()

    answerer = Answerer(s)
    answerer._openai = type("Client", (), {"chat": type("Chat", (), {"completions": Completions()})()})()
    text = "".join([t async for t in answerer.stream("How many points?", "", "")])
    assert text == "Four points." and seen[0]["model"] == "deepseek-flash"
    assert seen[0]["extra_body"] == {"thinking": {"type": "disabled"}}
    _ = [t async for t in answerer.stream("How many points?", "", "", deep=True)]
    assert seen[1]["model"] == "deepseek-v4-pro" and seen[1]["extra_body"]["thinking"]["type"] == "enabled"
    # the real client would be built with the DeepSeek key and URL
    answerer._openai = None
    client = answerer._client()
    assert str(client.base_url).startswith("https://api.deepseek.com") and client.api_key == "ds-key"


def test_dotenv_overrides_shell_and_last_duplicate_wins(tmp_path, monkeypatch):
    from classmate.config import load_dotenv, KEY_SOURCES
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-old-shell-key-0000")
    env = tmp_path / ".env"
    env.write_text('﻿DEEPSEEK_API_KEY=sk-first-key-1111\nexport DEEPSEEK_API_KEY = "sk-real​-key-ed22"\n')
    load_dotenv(str(env))
    import os
    assert os.environ["DEEPSEEK_API_KEY"] == "sk-real-key-ed22"
    assert KEY_SOURCES["DEEPSEEK_API_KEY"] == str(env)
    assert Settings(llm="deepseek").key("deepseek") == "sk-real-key-ed22"


# ---------------- screenshots ----------------
async def fake_grab(folder, region=False, max_side=1568):
    from PIL import Image
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"s{len(list(folder.iterdir()))}.jpg"
    Image.new("RGB", (40, 30), (200, 10, 10)).save(path)
    return path


async def test_tab_screenshot_is_answered_as_its_own_titled_question(tmp_path):
    session, ans = make_session(tmp_path)
    session.grab = fake_grab
    ans.gate.set()
    await say(session, "What is a kernel?")
    await settle()
    events = []
    orig = session.broadcast

    async def capture(e):
        events.append(e)
        await orig(e)
    session.broadcast = capture
    await session.ctrl(True)                 # an open speech capture must not swallow the screenshot
    await session.screenshot()
    await settle()
    await session.ctrl(False)
    session._last_shot = 0
    await session.screenshot(region=True)
    await settle()
    shots = [e for e in events if e["type"] == "question" and e["source"] == "screenshot"]
    assert [e["title"] for e in shots] == ["截图 1", "截图 2"]
    assert all(len(e["images"]) == 1 and e["images"][0].startswith("/shots/") for e in shots)
    image_calls = [c for c in ans.calls if c["images"]]
    assert len(image_calls) == 2 and image_calls[0]["q"] == ""
    assert image_calls[0]["history"][0][0] == "What is a kernel?"      # earlier answers still used
    # deep answer on a screenshot keeps its image
    await session.handle_command({"cmd": "deep", "id": shots[0]["id"]})
    await settle()
    assert ans.calls[-1]["images"] and session.titles[max(session.titles)] == "截图 1（深度）"


async def test_screenshot_debounce_and_served(tmp_path):
    s = Settings(session_dir=str(tmp_path))
    session = Session(s, lambda: None, asr=ScriptedASR(), answerer=FakeAnswerer())
    session.grab = fake_grab
    assert await session.screenshot() is not None
    assert await session.screenshot() is None            # window + global key in the same instant
    client = TestClient(TestServer(build_app(session)))
    await client.start_server()
    url = session.shot_url(next(session.shot_dir.iterdir()))
    resp = await client.get(url)
    assert resp.status == 200 and (await resp.read())[:2] == b"\xff\xd8"
    assert (await client.get("/shots/../../etc/passwd")).status == 404
    await client.close()


def test_summaries():
    from classmate.app import summarize
    long = summarize("so why do we use padding in convolution layers at all?")
    assert long.startswith("Why do we use padding") and long.endswith("…") and len(long) <= 36
    assert summarize("What is ReLU?") == "What is ReLU?"


async def test_image_content_and_ocr_paths(tmp_path, monkeypatch):
    from classmate import llm as llm_mod
    from classmate import screenshot as shot_mod
    img = await fake_grab(tmp_path)
    blocks, user = build_prompt("", "", "", has_images=True)
    assert "Extract the question or content and answer it fully" in user
    # Anthropic: image block before the text
    sent = {}

    class Stream:
        def __init__(self, **kw):
            sent.update(kw)
        async def __aenter__(self):
            async def gen():
                yield "B"
            self.text_stream = gen()
            return self
        async def __aexit__(self, *a):
            return False

    monkeypatch.setenv("ANTHROPIC_API_KEY", "x")
    a = llm_mod.Answerer(Settings(llm="anthropic"))
    a._anthropic = type("C", (), {"messages": type("M", (), {"stream": lambda self, **kw: Stream(**kw)})()})()
    out = "".join([t async for t in a.stream("", "", "", images=[img])])
    content = sent["messages"][0]["content"]
    assert out == "B" and content[0]["type"] == "image" and content[0]["source"]["media_type"] == "image/jpeg"
    assert content[-1]["type"] == "text"
    # DeepSeek: on-device OCR text instead of the picture
    monkeypatch.setattr(shot_mod, "ocr", lambda p, languages=None: "Q3. Which model? (b) homography")
    monkeypatch.setenv("DEEPSEEK_API_KEY", "k")
    ds = llm_mod.Answerer(load_settings(None, llm="deepseek"))
    seen = []

    class Completions:
        async def create(self, **kw):
            seen.append(kw)
            async def gen():
                yield type("Ch", (), {"choices": [type("C", (), {"delta": type("D", (), {"content": "b"})()})()]})()
            return gen()
    ds._openai = type("Client", (), {"chat": type("Chat", (), {"completions": Completions()})()})()
    _ = [t async for t in ds.stream("", "", "", images=[img])]
    user_msg = seen[0]["messages"][1]["content"]
    assert isinstance(user_msg, str) and "<screenshot_text>" in user_msg and "homography" in user_msg


async def test_ocr_text_shown_and_empty_ocr_explained(tmp_path, monkeypatch):
    from classmate import screenshot as shot_mod
    s = Settings(session_dir=str(tmp_path), llm="deepseek", vision="ocr")
    answerer = GatedAnswerer()
    answerer.gate.set()
    session = Session(s, lambda: None, asr=ScriptedASR(), answerer=answerer)
    session.grab = fake_grab
    events = []
    orig = session.broadcast

    async def capture(e):
        events.append(e)
        await orig(e)
    session.broadcast = capture
    monkeypatch.setattr(shot_mod, "ocr", lambda p, languages=None: "Q3. Which model? (b) homography")
    await session.screenshot()
    await settle()
    ocr_events = [e for e in events if e["type"] == "ocr"]
    assert ocr_events and "homography" in ocr_events[0]["text"]
    assert any(e["type"] == "ocr" for e in session.replay())
    monkeypatch.setattr(shot_mod, "ocr", lambda p, languages=None: "  ")
    session._last_shot = 0
    await session.screenshot()
    await settle()
    errors = [e for e in events if e["type"] == "error"]
    assert errors and "屏幕录制" in errors[-1]["message"]


def test_pick_window_skips_classmate_and_system_windows():
    from classmate.screenshot import pick_window
    big = {"Width": 1200, "Height": 800}
    windows = [
        {"kCGWindowNumber": 1, "kCGWindowLayer": 25, "kCGWindowOwnerName": "SystemUIServer", "kCGWindowBounds": big},
        {"kCGWindowNumber": 2, "kCGWindowLayer": 0, "kCGWindowOwnerPID": 999, "kCGWindowName": "Classmate",
         "kCGWindowOwnerName": "Python", "kCGWindowBounds": big},
        {"kCGWindowNumber": 3, "kCGWindowLayer": 0, "kCGWindowOwnerPID": 42, "kCGWindowName": "Classmate",
         "kCGWindowOwnerName": "Google Chrome", "kCGWindowBounds": big},          # browser-mode UI tab
        {"kCGWindowNumber": 4, "kCGWindowLayer": 0, "kCGWindowOwnerPID": 50, "kCGWindowOwnerName": "zoom.us",
         "kCGWindowBounds": {"Width": 120, "Height": 40}},                          # tiny toolbar
        {"kCGWindowNumber": 5, "kCGWindowLayer": 0, "kCGWindowOwnerPID": 50, "kCGWindowOwnerName": "zoom.us",
         "kCGWindowName": "Zoom Meeting", "kCGWindowBounds": big},
    ]
    assert pick_window(windows, own_pid=999) == 5
    assert pick_window(windows[:3], own_pid=999) is None


@pytest.mark.parametrize("text,filler", [("嗯。", True), ("是的。", True), ("好的", True), ("Okay.", True),
                                         ("Yeah", True), ("对对对。", True), ("Why?", False),
                                         ("你知道什么是 Redis 吗?", False), ("What is Redis", False),
                                         ("它就是截图，它识别不了。", False)])
def test_fillers(text, filler):
    from classmate.app import is_filler
    assert bool(is_filler(text)) is filler


async def test_every_turn_skips_fillers(tmp_path):
    session, ans = make_session(tmp_path, mode="every_turn")
    ans.gate.set()
    await say(session, "嗯。")
    await say(session, "Okay.")
    await say(session, "What is Redis used for?")
    await settle()
    assert [c["q"] for c in ans.calls] == ["What is Redis used for?"]


def test_screenshot_prompt_asks_for_extraction_and_full_answer():
    from classmate.llm import SCREENSHOT
    blocks, user = build_prompt("", "", "", has_images=True)
    assert blocks[0].startswith(SCREENSHOT) and "**题目：**" in blocks[0] and "up to 3 short" not in blocks[0]
    assert "answer it fully" in user
    blocks, _ = build_prompt("What is RANSAC?", "", "")
    assert "up to 3 short supporting bullet points" in blocks[0]
    _, user = build_prompt("", "", "", screen_text="Q1. 2+2=?")
    assert "<screenshot_text>" in user and "recognised from the screenshot" in user


async def test_screenshot_answers_get_long_token_budget(monkeypatch, tmp_path):
    from classmate.llm import Answerer
    monkeypatch.setenv("DEEPSEEK_API_KEY", "k")
    a = Answerer(load_settings(None, llm="deepseek", vision="ocr"))
    seen = []

    class Completions:
        async def create(self, **kw):
            seen.append(kw)
            async def gen():
                yield type("Ch", (), {"choices": [type("C", (), {"delta": type("D", (), {"content": "x"})()})()]})()
            return gen()
    a._openai = type("Client", (), {"chat": type("Chat", (), {"completions": Completions()})()})()
    img = await fake_grab(tmp_path)
    _ = [t async for t in a.stream("", "", "", images=[img], screen_text="Q1. what is 2+2")]
    _ = [t async for t in a.stream("What is 2+2?", "", "")]
    assert seen[0]["max_tokens"] == 3000 and seen[1]["max_tokens"] == 400


def test_auto_llm_and_notes_folder(monkeypatch, tmp_path):
    from classmate import config
    for env in ("ANTHROPIC_API_KEY", "DEEPSEEK_API_KEY", "OPENAI_API_KEY"):
        monkeypatch.delenv(env, raising=False)
    monkeypatch.setattr(config, "load_dotenv", lambda *a, **k: None)
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-x")
    s = config.load_settings(None)
    assert s.llm == "deepseek" and s.llm_model == "deepseek-flash"
    monkeypatch.setenv("ANTHROPIC_API_KEY", "a")
    assert config.load_settings(None).llm == "anthropic"
    assert config.load_settings(None, llm="deepseek").llm == "deepseek"     # explicit choice wins
    notes = tmp_path / "notes"
    notes.mkdir()
    monkeypatch.setattr(config, "NOTES_DIR", notes)
    assert config.load_settings(None).notes == [str(notes)]


# ---------------- Doubao / Qwen, settings panel, custom keys ----------------
def fake_openai(seen):
    class Completions:
        async def create(self, **kw):
            seen.append(kw)
            async def gen():
                yield type("Ch", (), {"choices": [type("C", (), {"delta": type("D", (), {"content": "x"})()})()]})()
            return gen()
    return type("Client", (), {"chat": type("Chat", (), {"completions": Completions()})()})()


@pytest.mark.parametrize("provider, env, url, fast, deep, off, on", [
    ("doubao", "ARK_API_KEY", "https://ark.cn-beijing.volces.com/api/v3", "doubao-seed-1-6-flash-250828",
     "doubao-seed-1-6-251015", {"thinking": {"type": "disabled"}}, {"thinking": {"type": "enabled"}}),
    ("qwen", "DASHSCOPE_API_KEY", "https://dashscope.aliyuncs.com/compatible-mode/v1", "qwen-flash", "qwen-plus",
     {"enable_thinking": False}, {"enable_thinking": True}),
])
async def test_doubao_and_qwen_providers(monkeypatch, provider, env, url, fast, deep, off, on):
    from classmate.llm import Answerer
    monkeypatch.setenv(env, "k-123")
    s = load_settings(None, llm=provider)
    assert (s.llm_model, s.llm_deep_model, s.llm_base_url) == (fast, deep, url)
    a, seen = Answerer(s), []
    a._openai = fake_openai(seen)
    _ = [t async for t in a.stream("Why?", "", "")]
    _ = [t async for t in a.stream("Why?", "", "", deep=True)]
    assert seen[0]["model"] == fast and seen[0]["extra_body"] == off
    assert seen[1]["model"] == deep and seen[1]["extra_body"] == on
    a._openai = None
    client = a._client()
    assert str(client.base_url).rstrip("/") == url and client.api_key == "k-123"


async def test_qwen_screenshots_use_the_vision_model(monkeypatch, tmp_path):
    from classmate.llm import Answerer
    monkeypatch.setenv("DASHSCOPE_API_KEY", "k")
    a, seen = Answerer(load_settings(None, llm="qwen")), []
    a._openai = fake_openai(seen)
    img = await fake_grab(tmp_path)
    _ = [t async for t in a.stream("", "", "", images=[img])]
    assert seen[0]["model"] == "qwen3-vl-plus"
    assert seen[0]["messages"][1]["content"][0]["type"] == "image_url"


def test_auto_llm_picks_doubao_or_qwen_from_keys(monkeypatch):
    monkeypatch.setenv("DASHSCOPE_API_KEY", "k")
    assert load_settings(None).llm == "qwen"
    monkeypatch.setenv("ARK_API_KEY", "k")
    assert load_settings(None).llm == "doubao"


def test_key_validation():
    from classmate.config import DEFAULT_KEYS, normalize_key, validate_keys
    assert normalize_key("Shift+ArrowUp") == "shift+up" and normalize_key("cmd+shift+K") == "shift+cmd+k"
    assert validate_keys({}) == DEFAULT_KEYS
    assert validate_keys({"screenshot": "F5", "capture": "Option"})["capture"] == "alt"
    with pytest.raises(ValueError, match="冲突"):
        validate_keys({"deep": "c"})                         # copy is already C
    with pytest.raises(ValueError, match="冲突"):
        validate_keys({"screenshot": "d", "deep": "shift+d"})   # region defaults to Shift + screenshot
    combo = validate_keys({"screenshot": "Control+Option+S", "capture": "alt+space"})
    assert (combo["screenshot"], combo["region"], combo["capture"]) == ("ctrl+alt+s", "ctrl+alt+shift+s", "alt+space")
    assert validate_keys({"screenshot": "cmd+shift+2"})["region"] == ""        # no free Shift variant
    assert validate_keys({"screenshot": "f5", "region": "f6"})["region"] == "f6"
    with pytest.raises(ValueError, match="开始收录"):
        validate_keys({"screenshot": "ctrl+s"})             # holding Ctrl already starts a capture
    with pytest.raises(ValueError):
        validate_keys({"screenshot": "ctrl"})
    with pytest.raises(ValueError):
        validate_keys({"mode": "up"})                       # arrows move the window
    with pytest.raises(ValueError):
        validate_keys({"nope": "x"})
    s = Settings(screenshot_key="f5")                        # old config files still work
    assert s.keys["screenshot"] == "f5" and s.keys["region"] == "shift+f5"


def test_global_shortcut_matching():
    from enum import Enum
    from classmate.desktop import Shortcut

    class Key(Enum):
        ctrl = 1; ctrl_l = 2; alt = 3; alt_l = 4; shift = 5; shift_l = 6; cmd = 7; tab = 8; space = 9; esc = 10

    class KeyCode:
        def __init__(self, char=None, vk=None):
            self.char, self.vk = char, vk
    kb = type("Keyboard", (), {"Key": Key, "KeyCode": KeyCode})
    shot = Shortcut(kb, "ctrl+alt+s")
    assert shot.pressed(KeyCode("ß", 1), {"ctrl", "alt"})          # ⌥ changes the character, not the key code
    assert not shot.pressed(KeyCode("s", 1), {"ctrl"}) and not shot.pressed(KeyCode("s", 1), {"ctrl", "alt", "shift"})
    assert Shortcut(kb, "tab").pressed(Key.tab, set()) and not Shortcut(kb, "tab").pressed(Key.tab, {"shift"})
    assert Shortcut(kb, "ctrl").pressed(Key.ctrl_l, set())
    assert Shortcut(kb, "alt+space").pressed(Key.space, {"alt"})


async def test_settings_panel_saves_and_applies(tmp_path, monkeypatch):
    from classmate import config
    from classmate.llm import Answerer
    s = Settings(session_dir=str(tmp_path / "sessions"), llm="deepseek")
    session = Session(s, lambda: None, asr=ScriptedASR())
    events, changed = [], []

    async def capture(e):
        events.append(e)
    session.broadcast = capture
    session.on_keys_changed = lambda: changed.append(True)
    monkeypatch.setattr(Answerer, "warmup", lambda self: asyncio.sleep(0))
    await session.handle_command({"cmd": "get_config"})
    cfg = events[-1]
    assert cfg["type"] == "config" and "doubao" in cfg["providers"] and cfg["values"]["llm"] == "deepseek"
    assert cfg["api_keys"]["ARK_API_KEY"] == ""
    await session.handle_command({"cmd": "save_config",
                                  "values": {"llm": "doubao", "llm_model": "", "keys": {"screenshot": "f5"}},
                                  "api_keys": {"ARK_API_KEY": " ark-secret-key-1234 "}})
    assert s.llm == "doubao" and s.llm_model == "doubao-seed-1-6-flash-250828" and s.keys["screenshot"] == "f5"
    assert session.answerer.s.llm == "doubao" and changed == [True]
    assert "ARK_API_KEY=ark-secret-key-1234" in (config.APP_DIR / ".env").read_text()
    shown = [e for e in events if e["type"] == "config"][-1]["api_keys"]["ARK_API_KEY"]
    assert "secret" not in shown and shown.endswith("(19 chars)")       # the page only sees a masked key
    assert {e["type"] for e in events} >= {"status", "config", "config_saved"}
    # remembered for the next start, above config.yaml, below command-line flags
    (tmp_path / "c.yaml").write_text("llm: deepseek\nllm_model: deepseek-flash\n")
    again = load_settings(str(tmp_path / "c.yaml"))
    assert (again.llm, again.llm_model, again.keys["screenshot"]) == ("doubao", "doubao-seed-1-6-flash-250828", "f5")
    assert load_settings(str(tmp_path / "c.yaml"), llm="qwen").llm_model == "qwen-flash"
    saved = config.read_ui_settings()
    assert saved["llm_model"] is None and saved["keys"] == {"screenshot": "f5", "region": "shift+f5"}
    with pytest.raises(ValueError):
        await session.save_config({"keys": {"deep": "c"}})
    assert s.keys["deep"] == "d"


def test_save_env_key_replaces_existing_lines(tmp_path, monkeypatch):
    from classmate import config
    env = tmp_path / ".env"
    env.write_text("# keys\nARK_API_KEY=old\nexport ARK_API_KEY=older\nDEEPSEEK_API_KEY=ds\n")
    monkeypatch.setattr(config, "ENV_FILE", env)
    config.save_env_key("ARK_API_KEY", "new")
    assert env.read_text() == "# keys\nARK_API_KEY=new\nDEEPSEEK_API_KEY=ds\n"
    import os
    assert os.environ["ARK_API_KEY"] == "new" and oct(env.stat().st_mode)[-3:] == "600"


async def test_missing_asr_key_does_not_start(tmp_path):
    s = Settings(session_dir=str(tmp_path), llm="deepseek")
    session = Session(s, lambda: None)
    events = []

    async def capture(e):
        events.append(e)
    session.broadcast = capture
    assert set(session.missing_keys()) == {"ASSEMBLYAI_API_KEY", "DEEPSEEK_API_KEY"}
    await session.start()
    assert not session.listening and "ASSEMBLYAI_API_KEY" in events[-1]["message"]


async def test_opacity_slider_sets_window_and_is_remembered(tmp_path):
    from classmate import config

    class Api:
        def __init__(self):
            self.values = []
        def set_opacity(self, v):
            self.values.append(v)
        def move(self, dx, dy):
            pass
    api = Api()
    session = Session(Settings(session_dir=str(tmp_path)), lambda: None, asr=ScriptedASR(), answerer=FakeAnswerer())
    async with TestClient(TestServer(build_app(session, api))) as client:
        ws = await client.ws_connect("/ws")
        first = json.loads((await ws.receive()).data)
        assert first["native"] is True and first["opacity"] == 1.0
        await ws.send_json({"cmd": "opacity", "value": 0.55})
        await ws.send_json({"cmd": "opacity", "value": 0.05, "save": True})      # never fully invisible
        for _ in range(50):
            await asyncio.sleep(0.01)
            if config.read_ui_settings().get("opacity"):
                break
        await ws.close()
    assert api.values == [0.55, 0.2] and config.read_ui_settings()["opacity"] == 0.2
    assert load_settings(None).opacity == 0.2


def test_answers_are_always_chinese():
    from classmate.llm import looks_english
    for args in (("What is a homography?", "", ""), ("", "", "")):
        blocks, user = build_prompt(*args, has_images=not args[0])
        assert blocks[0].startswith("必须全程使用简体中文作答") and user.startswith("【请用简体中文回答")
        assert user.endswith("请用简体中文回答。")
    _, strict = build_prompt("Why?", "", "", strict=True)
    assert strict.startswith("【最高优先级】")
    english = "The ratio test rejects ambiguous matches: keep a match only if the best distance is clearly smaller."
    assert looks_english(english) and not looks_english("比值检验（ratio test）用来剔除模糊的匹配，只保留最近邻明显更近的匹配。")
    assert not looks_english("答案如下：\n```python\n" + "x = compute_homography(points)\n" * 5 + "```")


async def test_english_answer_is_replaced_by_a_chinese_one(monkeypatch, tmp_path):
    from classmate.llm import RESTART, Answerer
    monkeypatch.setenv("DEEPSEEK_API_KEY", "k")
    a = Answerer(load_settings(None, llm="deepseek"))
    calls = []

    async def once(*args, strict=False):
        calls.append(strict)
        text = ("单应性（homography）描述两个平面之间的投影变换。" * 2 if strict else
                "A homography maps points between two planes; it has eight degrees of freedom, so we need four "
                "point correspondences to estimate it.")
        for word in text.split(" "):
            yield word + " "
    a._stream_once = once
    out = [t async for t in a.stream("What is a homography?", "", "")]
    assert calls == [False, True] and RESTART in out
    assert "单应性" in "".join(t for t in out[out.index(RESTART) + 1:])
    # the session clears the English text in the UI and keeps only the Chinese answer
    session = Session(Settings(session_dir=str(tmp_path)), lambda: None, asr=ScriptedASR(), answerer=a)
    events = []

    async def capture(e):
        events.append(e)
    session.broadcast = capture
    qid = session.ask("What is a homography?", source="typed")
    await session._answers[qid]
    assert any(e["type"] == "answer_reset" for e in events) and session.answers[qid].startswith("单应性")
    replay = "".join(e["text"] for e in session.replay() if e["type"] == "answer_delta")
    assert "homography maps" not in replay and replay.startswith("单应性")


async def test_toggle_ctrl_records_then_answers_on_second_tap(tmp_path):
    session, ans = make_session(tmp_path, style="toggle")
    ans.gate.set()
    session.listening = True                     # pretend the microphone is already running
    await session.ctrl(True)                     # key down does nothing in toggle style
    await session.ctrl(False, tap=True)          # first tap: start recording
    assert session.hold
    await say(session, "What is the receptive field")
    await session.on_transcript(Transcript("and how does it grow", final=False))
    assert not ans.calls                         # nothing answered while recording
    session._last_tap = 0
    await session.ctrl(False, tap=True)          # second tap: answer now, live words included
    await settle()
    assert not session.hold and ans.calls[0]["q"] == "What is the receptive field and how does it grow"
    assert not session.listening                 # microphone off until the next tap
    assert "and how does it grow" in session.transcript.window(5)    # live words kept as context
    session._last_tap = 0
    await session.ctrl(False, tap=False)         # Ctrl+C in another app: not a tap
    assert not session.hold
    session.start = lambda: setattr(session, "listening", True) or asyncio.sleep(0)   # no real microphone
    await session.ctrl(False, tap=True)
    await session.ctrl(False, tap=True)          # same tap from window + global listener: counted once
    assert session.hold and session.listening
    await session.cancel_capture()               # Esc
    assert not session.hold and not session.cap and not session.listening


def test_global_listener_ignores_system_defined_events():
    """Caps Lock (中/英 switch) sends a system-defined event; turning it into an NSEvent off the main
    thread crashed the app, so the listener must neither subscribe to it nor handle it."""
    import sys
    if sys.platform != "darwin":
        pytest.skip("macOS only")
    import Quartz
    from pynput import keyboard
    from classmate.desktop import NS_SYSTEM_DEFINED, listener_class
    cls = listener_class(keyboard)
    assert not cls._EVENTS & Quartz.CGEventMaskBit(NS_SYSTEM_DEFINED)
    assert cls._EVENTS & Quartz.CGEventMaskBit(Quartz.kCGEventFlagsChanged)      # Ctrl still works
    seen = []
    listener = cls(on_press=seen.append, on_release=seen.append)
    listener._handle_message(None, NS_SYSTEM_DEFINED, object(), None, False)   # would crash in NSEvent
    assert seen == []


async def test_preferred_code_language(monkeypatch, tmp_path):
    from classmate.llm import Answerer
    blocks, _ = build_prompt("", "", "", has_images=True, code_language="Java")
    assert "in Java" in blocks[0] and "already given" in blocks[0]
    assert "Code:" not in build_prompt("", "", "", has_images=True)[0][0]      # empty = follow the question
    monkeypatch.setenv("DEEPSEEK_API_KEY", "k")
    a, seen = Answerer(load_settings(None, llm="deepseek", code_language="C++")), []
    a._openai = fake_openai(seen)
    _ = [t async for t in a._stream_once("Write a function to reverse a list", "", "")]
    assert "in C++" in seen[0]["messages"][0]["content"]


async def test_append_key_adds_to_the_current_question(tmp_path):
    session, ans = make_session(tmp_path, style="toggle")
    ans.gate.set()
    session.start = lambda: setattr(session, "listening", True) or asyncio.sleep(0)   # no real microphone
    events = []
    orig = session.broadcast

    async def capture(e):
        events.append(e)
        await orig(e)
    session.broadcast = capture
    await session.ctrl(False, tap=True)
    await say(session, "What is a homography")
    session._last_tap = 0
    await session.ctrl(False, tap=True)
    await settle()
    qid = session.last_qid
    session._last_tap = 0
    await session.append_tap()                   # ` : record more for Q1
    assert session.hold and session.cap_target == qid
    await say(session, "and how many points do we need")
    session._last_tap = 0
    await session.append_tap()                   # ` again: regenerate Q1 with the addition
    await settle()
    assert ans.calls[-1]["q"] == "What is a homography and how many points do we need"
    assert session.last_qid == qid and len(session.questions) == 1
    assert any(e["type"] == "question" and e["replace"] for e in events)    # UI marks it "已补充"
    assert not session.listening
    session._last_tap = 0
    await session.append_tap()
    session._last_tap = 0
    await session.ctrl(False, tap=True)          # Ctrl also ends it; nothing new said: no regeneration
    await settle()
    assert len(ans.calls) == 2 and not session.cap


def test_append_key_default_and_matching():
    from enum import Enum
    from classmate.config import DEFAULT_KEYS
    from classmate.desktop import Shortcut
    assert DEFAULT_KEYS["append"] == "`" and Settings().keys["append"] == "`"

    class Key(Enum):
        shift = 1

    class KeyCode:
        def __init__(self, char=None, vk=None):
            self.char, self.vk = char, vk
    kb = type("Keyboard", (), {"Key": Key, "KeyCode": KeyCode})
    assert Shortcut(kb, "`").pressed(KeyCode("·", 50), set())           # Chinese input method types ·


# ---------------- authorization codes ----------------
@pytest.fixture
def test_keys(monkeypatch):
    """A throwaway signing key instead of the owner's (whose master code is never in the repo)."""
    import base64
    from cryptography.hazmat.primitives import serialization as ser
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    key = Ed25519PrivateKey.generate()
    pub = key.public_key().public_bytes(ser.Encoding.Raw, ser.PublicFormat.Raw)
    monkeypatch.setattr(licensing, "PUBLIC_KEY", base64.b64encode(pub).decode())
    monkeypatch.setattr(licensing, "load", REAL_LICENSE_LOAD)
    seed = key.private_bytes(ser.Encoding.Raw, ser.PrivateFormat.Raw, ser.NoEncryption())
    return "CM-MASTER-" + base64.b32encode(seed).decode().rstrip("=")


def test_master_code_issues_expiring_codes(test_keys):
    master = licensing.parse(test_keys.lower())                 # case and spaces do not matter
    assert master.kind == "master" and master.expires is None and master.can_issue
    now = 1_800_000_000
    code = licensing.issue(master, 7, now=now)
    user = licensing.parse(code, now=now + 6 * 86400)
    assert user.kind == "user" and not user.can_issue and user.expires == now + 7 * 86400
    with pytest.raises(ValueError, match="过期"):
        licensing.parse(code, now=now + 8 * 86400)
    with pytest.raises(PermissionError):
        licensing.issue(user, 7)                                 # issued codes cannot issue codes
    tampered = code[:-3] + ("AAA" if code[-3:] != "AAA" else "BBB")
    with pytest.raises(ValueError):
        licensing.parse(tampered, now=now)
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    from cryptography.hazmat.primitives import serialization as ser
    import base64
    other = Ed25519PrivateKey.generate().private_bytes(ser.Encoding.Raw, ser.PrivateFormat.Raw, ser.NoEncryption())
    with pytest.raises(ValueError):
        licensing.parse("CM-MASTER-" + base64.b32encode(other).decode())   # someone else's key
    with pytest.raises(ValueError, match="CM-"):
        licensing.parse("hello")


def test_code_is_remembered_and_clock_rollback_is_caught(test_keys):
    master = licensing.parse(test_keys)
    now = time_now = 1_800_000_000
    lic = licensing.parse(licensing.issue(master, 30, now=now), now=now)
    licensing.save(lic)
    licensing._write(lic.code, now + 10 * 86400)                 # last seen 10 days later...
    assert licensing.load(now=now + 2 * 86400) is None           # ...then the clock went back
    assert licensing.load(now=now + 11 * 86400).kind == "user"
    licensing.forget()
    assert licensing.load(now=time_now) is None


async def test_app_requires_a_code_and_only_master_issues(test_keys, tmp_path):
    session = Session(Settings(session_dir=str(tmp_path)), lambda: None, asr=ScriptedASR(), answerer=FakeAnswerer())
    assert session.license is None
    async with TestClient(TestServer(build_app(session))) as client:
        ws = await client.ws_connect("/ws")
        assert json.loads((await ws.receive()).data)["license"] is None

        async def until(kind, check=lambda m: True):
            while True:
                m = json.loads((await ws.receive()).data)
                if m["type"] == kind and check(m):
                    return m
        await ws.send_json({"cmd": "ask", "text": "What is RANSAC?"})
        assert "授权码" in (await until("activate_failed"))["message"]
        await ws.send_json({"cmd": "activate", "code": "CM-WRONG"})
        assert "无效" in (await until("activate_failed"))["message"]
        await ws.send_json({"cmd": "activate", "code": test_keys})
        assert (await until("status"))["license"] == {"kind": "master", "expires": None, "can_issue": True}
        await ws.send_json({"cmd": "issue_code", "days": 3})
        code = (await until("issued_code"))["code"]
        await ws.send_json({"cmd": "logout"})
        assert (await until("status"))["license"] is None
        await ws.send_json({"cmd": "activate", "code": code})
        info = (await until("status", lambda m: m["license"]))["license"]
        assert info["kind"] == "user" and not info["can_issue"]
        await ws.send_json({"cmd": "issue_code", "days": 3})
        assert "主授权码" in (await until("error"))["message"]
        await ws.close()
    from classmate import config
    assert json.loads((config.APP_DIR / "license.txt").read_text())["code"] == code     # remembered
